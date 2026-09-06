"""Word。段落スタイルから見出しレベルを取る。ページ番号は取得不能なので None。"""

from __future__ import annotations

import io
import re

from app.core.parsers.base import Block, ParsedDocument, ParseError

#: 見出しスタイル名。日本語版 Word は「見出し 1」、英語版は "Heading 1" になる。
HEADING_STYLE = re.compile(r"(?:Heading|見出し)\s*(\d+)")


def parse(data: bytes) -> ParsedDocument:
    """Word から段落と表を Block にする。

    ページ番号は python-docx では取得できない (改ページは印刷時に決まるため)。
    推測はせず None のままにし、出力時に「不明」と表示する。
    """
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ParseError("python-docx が未インストールです: pip install python-docx") from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Word文書を読み込めません: {exc}") from exc

    blocks: list[Block] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        level: int | None = None
        style_name = (para.style.name if para.style is not None else "") or ""
        m = HEADING_STYLE.search(style_name)
        if m:
            level = int(m.group(1))
        # 表題は章より上位だが、階層としては最上位と同じ扱いで足りる
        elif style_name in ("Title", "表題"):
            level = 1
        blocks.append(Block(text=text, heading_level=level))

    # 表は行単位で1ブロックにする。セルを連結するのは行内だけで、行はまたがない
    for t_index, table in enumerate(document.tables, start=1):
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            joined = " / ".join(c for c in cells if c)
            if joined:
                blocks.append(Block(text=joined, is_table=True, locator=f"表{t_index}"))

    if not blocks:
        raise ParseError("Word文書からテキストを抽出できませんでした。")
    return ParsedDocument(blocks=blocks, meta={"format": "docx"})
