"""Excel。

「No / 章 / 節 / 分類 / 規定内容 / 区分 / 重要度」のようなヘッダを持つ表であれば、
列の役割を判定して規定文だけを取り出し、章・節・区分・重要度は標準書の記載として
そのまま引き渡す。ヘッダを判定できない場合は、行全体を1ブロックとして扱う。

所在はシート名と行番号で表す (ページ番号は存在しないので推測しない)。
"""

from __future__ import annotations

import io

from app.core.parsers.base import Block, ParsedDocument, ParseError
from app.core.parsers.table_schema import (
    detect_tables,
    map_chapter,
    map_rule_type,
    map_section,
    map_severity,
)

#: 1シートあたりに読む最大行数 (巨大ファイルで固まらないようにする)
MAX_ROWS_PER_SHEET = 20000


def parse(data: bytes) -> ParsedDocument:
    """Excel を Block にする。

    ヘッダを判定できた表は列ごとに意味を持たせて取り出し、判定できなければ
    行全体を1ブロックとして扱う。後者に落ちても規定文は失われない。
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ParseError("openpyxl が未インストールです: pip install openpyxl") from exc

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ParseError(f"Excelを読み込めません: {exc}") from exc

    blocks: list[Block] = []
    structured_sheets = 0
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append([("" if v is None else str(v).strip()) for v in row])
            if len(rows) >= MAX_ROWS_PER_SHEET:
                break

        blocks.append(Block(text=str(ws.title), heading_level=1, locator=str(ws.title)))

        tables = detect_tables(rows)
        if tables:
            structured_sheets += 1
            # 表ごとに、その表のヘッダで読む。1枚に複数の表が並ぶことがある
            for table in tables:
                blocks.extend(_structured_blocks(ws.title, rows, table))
        else:
            blocks.extend(_flat_blocks(ws.title, rows))

    wb.close()
    if len(blocks) <= len(wb.worksheets):
        raise ParseError("Excelからテキストを抽出できませんでした。")
    return ParsedDocument(
        blocks=blocks,
        meta={"format": "xlsx", "structured_sheets": str(structured_sheets)},
    )


def _structured_blocks(sheet: str, rows: list[list[str]], table) -> list[Block]:
    """1つの表から、規定内容の列だけを本文として取り出す。

    章・節・区分・重要度は標準書がその列に書いている値なので、推測ではなく
    記載そのものとして hint に載せて後段へ渡す。

    読むのは table が示す行範囲だけ。シート全体に当てると、次の表の行を
    この表の列対応で読んでしまう。
    """
    header_index, schema = table.header_index, table.schema
    rule_column = schema.rule_column
    assert rule_column is not None  # detect_tables が保証する

    def cell(row: list[str], role: str) -> str:
        index = schema.column_of(role)
        if index is None or index >= len(row):
            return ""
        return row[index]

    out: list[Block] = []
    body = rows[header_index + 1 : table.end_index + 1]
    for offset, row in enumerate(body, start=header_index + 2):
        if rule_column >= len(row):
            continue
        text = row[rule_column].strip()
        if not text:
            continue
        out.append(
            Block(
                text=text,
                is_table=True,
                locator=f"{sheet}!{offset}行",
                chapter=map_chapter(cell(row, "chapter")) or None,
                section=map_section(cell(row, "section")) or None,
                category_hint=cell(row, "category") or None,
                rule_type_hint=map_rule_type(cell(row, "rule_type")),
                severity_hint=map_severity(cell(row, "severity")),
                note_hint=cell(row, "note") or None,
            )
        )
    return out


def _flat_blocks(sheet: str, rows: list[list[str]]) -> list[Block]:
    """ヘッダを判定できないシート。行を連結した従来どおりの扱い。"""
    out: list[Block] = []
    for row_index, row in enumerate(rows, start=1):
        cells = [c for c in row if c]
        if not cells:
            continue
        out.append(
            Block(
                text=" / ".join(cells),
                is_table=True,
                locator=f"{sheet}!{row_index}行",
            )
        )
    return out
